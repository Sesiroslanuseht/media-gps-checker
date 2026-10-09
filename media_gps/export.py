from __future__ import annotations

import os
import re
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from .core import COMPLETE, MISSING, INVALID, UNKNOWN

HEADERS = ['序号', '路径', '文件类型', '媒体类别', '拍摄时间（元数据原值）', '纬度', '经度', 'GPS 状态', '原因', '动态部分', '动态/配对依据', '配对成员路径', '成员数']


def clean(value):
    if not isinstance(value, str):
        return value
    # Excel rejects these controls. Force string cell type to prevent formula injection.
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '�', value)[:32767]


def append(ws, values):
    cells = []
    for value in values:
        cell = WriteOnlyCell(ws, value=clean(value))
        if isinstance(value, str):
            cell.data_type = 's'
        cells.append(cell)
    ws.append(cells)


def export_xlsx(result, destination):
    dest = Path(destination).absolute()
    if dest.suffix.lower() != '.xlsx':
        raise ValueError('导出文件必须以 .xlsx 结尾')
    # Exclusive creation: never overwrite any existing file, including an original file.
    # On failure remove only the file created by this operation.
    with dest.open('xb') as target:
        try:
            book = Workbook(write_only=True)
            summary = book.create_sheet('扫描说明')
            for row in [
                ['项目', '内容'], ['生成时间', datetime.now().isoformat(timespec='seconds')],
                ['扫描目录', result.root], ['范围', result.scope], ['扫描状态', '已取消（部分结果）' if result.cancelled else '已完成'],
                ['发现候选文件', result.discovered], ['已读取物理文件', len(result.records)],
                ['所选范围媒体条目（合并确认配对）', len(result.assets)], ['非候选文件跳过数', result.skipped],
                ['ExifTool 版本', result.exiftool_version],
                ['说明', 'GPS 缺失仅指受支持元数据字段未发现经纬度，不能证明传输导致丢失。异常包括 0,0 待核实。'],
                ['动态照片', '结构检查不保证视频可播放；不保证识别所有小米型号、私有格式或经过编辑的动态照片。'],
                ['配对 GPS', '各成员单独检查；主表最严重状态优先：异常 > 无法判断 > 缺失 > 完整。不以视频 GPS 补全照片。'],
                ['视频 GPS', '读取容器标准元数据；不解析每帧轨迹或所有私有 GPS 数据。'],
                ['时间', '保留元数据原值，不自行转换时区；无拍摄时间则留空，不使用文件修改时间冒充。'],
            ]:
                append(summary, row)
            for state in (COMPLETE, MISSING, INVALID, UNKNOWN):
                append(summary, [state, sum(a.gps_status == state for a in result.assets)])
            for notice in result.notices:
                append(summary, ['提示', notice])
            main = book.create_sheet('媒体结果')
            main.freeze_panes = 'A2'
            append(main, HEADERS)
            detail = book.create_sheet('成员明细')
            detail.freeze_panes = 'A2'
            append(detail, ['媒体序号', '路径', '文件类型', '拍摄时间', '纬度', '经度', 'GPS 状态', '原因', 'GPS 原始字段', '动态状态', '动态依据'])
            for i, asset in enumerate(result.assets, 1):
                r = asset.members[0]
                append(main, [i, r.path, r.file_type, asset.kind, r.taken, r.gps.latitude, r.gps.longitude, asset.gps_status,
                              '; '.join(f'{Path(m.path).name}: {m.gps.reason}' for m in asset.members), asset.motion,
                              asset.evidence, '\n'.join(m.path for m in asset.members[1:]), len(asset.members)])
                for m in asset.members:
                    append(detail, [i, m.path, m.file_type, m.taken, m.gps.latitude, m.gps.longitude, m.gps.status,
                                    m.gps.reason, m.gps.evidence, m.motion, m.evidence])
            main.auto_filter.ref = f'A1:M{len(result.assets)+1}'
            book.save(target)
        except BaseException:
            target.close()
            dest.unlink(missing_ok=True)
            raise
    return str(dest)
