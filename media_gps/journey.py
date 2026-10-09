"""One KML per capture day, plus an ordered workbook. Source media are never opened."""
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import math
import re
import shutil
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from .core import COMPLETE, PHOTO, MOTION, VIDEO, vals
from .capture_time import CaptureTime, capture_time, fixed_timezone
from .export import append

KML = 'http://www.opengis.net/kml/2.2'
DESCRIPTION = '拍照/录像地点按拍摄时间连接；不是连续记录的真实步行或行车轨迹。坐标按 WGS84，不上传照片。'


@dataclass
class JourneyEntry:
    asset: object
    time: CaptureTime
    reason: str = ''


def prepare_journeys(result, offset='+08:00'):
    fixed_timezone(offset)
    entries = []
    for asset in result.assets:
        r = asset.members[0]
        t = capture_time(r, offset)
        problems = []
        if asset.kind not in (PHOTO, MOTION, VIDEO):
            problems.append('未确认的媒体类型')
        # A paired motion photo contributes one point, always from its primary still.
        if r.gps.status != COMPLETE:
            problems.append(r.gps.status + '：' + r.gps.reason)
        else:
            try:
                a, b = float(r.gps.latitude), float(r.gps.longitude)
                if not (math.isfinite(a) and math.isfinite(b) and abs(a) <= 90 and abs(b) <= 180) or (a == 0 and b == 0):
                    raise ValueError()
            except (ValueError, TypeError):
                problems.append('坐标数值不可用')
        for datum in vals(r.metadata, 'GPSMapDatum'):
            if re.sub(r'[^a-z0-9]', '', str(datum).lower()) not in ('wgs84', 'wgs1984'):
                problems.append(f'坐标基准不是已确认的 WGS84：{datum}')
        if t.utc is None:
            problems.append(t.note)
        entries.append(JourneyEntry(asset, t, '; '.join(problems)))
    entries.sort(key=lambda e: (e.time.utc is None, e.time.utc.timestamp() if e.time.utc else 0,
                               e.asset.members[0].path.casefold(), e.asset.members[0].path))
    days = {}
    for entry in entries:
        if not entry.reason:
            days.setdefault(entry.time.local.date().isoformat(), []).append(entry)
    return entries, days


def xml_text(value):
    # XML 1.0 text and HTML description safety. ElementTree escapes markup.
    return ''.join(c for c in str(value) if c in '\t\n\r' or 0x20 <= ord(c) <= 0xD7FF or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF)


def element(parent, name, value=None, **attrs):
    e = ET.SubElement(parent, name, attrs)
    if value is not None:
        e.text = xml_text(value)
    return e


def utc_text(dt):
    return dt.isoformat().replace('+00:00', 'Z')


def coordinates(record):
    # KML order is longitude, latitude, altitude. No fake elevation: use ground clamp.
    return f'{float(record.gps.longitude):.8f},{float(record.gps.latitude):.8f},0'


def kml_bytes(day, entries, partial=False):
    import html
    root = ET.Element('kml', {'xmlns': KML})
    doc = element(root, 'Document')
    element(doc, 'name', day + '行程' + ('（扫描未完成）' if partial else ''))
    element(doc, 'description', DESCRIPTION)
    style = element(doc, 'Style', id='route')
    line = element(style, 'LineStyle')
    element(line, 'color', 'ffdd6600')
    element(line, 'width', '3')
    if len(entries) >= 2:
        line = element(doc, 'Placemark')
        element(line, 'name', day + ' 拍摄地点连线')
        element(line, 'description', DESCRIPTION)
        element(line, 'styleUrl', '#route')
        geom = element(line, 'LineString')
        element(geom, 'tessellate', '1')
        element(geom, 'altitudeMode', 'clampToGround')
        element(geom, 'coordinates', '\n'.join(coordinates(e.asset.members[0]) for e in entries))
    for i, e in enumerate(entries, 1):
        r = e.asset.members[0]
        mark = element(doc, 'Placemark')
        element(mark, 'name', f'{i:03d} {e.time.local.strftime("%H:%M:%S")} {Path(r.path).name}')
        text = f'拍摄时间：{e.time.local.isoformat()}\n文件：{Path(r.path).name}\n类型：{e.asset.kind}\n{e.time.note}'
        if len(e.asset.members) > 1:
            text += '\n动态照片按主照片定位，配对视频不重复计点'
        element(mark, 'description', html.escape(text))  # KML descriptions may be rendered as HTML.
        ts = element(mark, 'TimeStamp')
        element(ts, 'when', utc_text(e.time.utc))
        point = element(mark, 'Point')
        element(point, 'altitudeMode', 'clampToGround')
        element(point, 'coordinates', coordinates(r))
    # Avoid global register_namespace races between concurrent exports.
    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


def write_workbook(result, entries, days, offset, path):
    book = Workbook(write_only=True)
    info = book.create_sheet('行程说明')
    for row in [
        ['项目', '内容'], ['扫描目录', result.root], ['范围', result.scope],
        ['扫描状态', '已取消，只含部分结果' if result.cancelled else '已完成'],
        ['行程时区', 'UTC' + offset], ['行程天数', len(days)],
        ['进入轨迹的点数', sum(len(v) for v in days.values())], ['跳过条目', sum(bool(e.reason) for e in entries)],
        ['顺序', '按拍摄时间先后；同时刻按文件路径稳定排序，不能判断同时刻内真实先后。'],
        ['日期', '元数据带时区先转换到行程时区，再按日期分组；没时区的照片按行程时区解释。'],
        ['视频时间', 'QuickTime 整数日期按 UTC 标准解释；部分设备写入本地时间，需人工核对。'],
        ['动态照片', '只使用确认配对条目的主照片坐标和时间，不向配对视频借用数据。'],
        ['坐标', '按 WGS84 输出；缺少基准字段时假设为 WGS84，不转换为 GCJ-02/BD-09。'],
        ['地图说明', DESCRIPTION],
        ['无记录路段', '同日有效点之间直线连接，未记录的路段及跳过文件可能形成长连线；不能据此计算真实路程。'],
        ['隐私', 'KML 仅包含文件名、时间和坐标，无照片或完整路径；Excel 包含源文件完整路径。'],
    ]:
        append(info, row)
    for notice in result.notices:
        append(info, ['扫描提示', notice])
    headers = ['排序序号', '行程日期', '拍摄时间（行程时区）', '拍摄时间（UTC）', '原始时间', '时间来源', '时区/时间说明', '主文件路径', '媒体类型', '纬度', '经度', '主文件 GPS 状态', '轨迹处理', '跳过原因', '动态部分', '配对成员']
    all_rows = book.create_sheet('拍摄顺序清单')
    skipped = book.create_sheet('跳过项')
    for sheet in (all_rows, skipped):
        sheet.freeze_panes = 'A2'
        append(sheet, headers)
    for i, e in enumerate(entries, 1):
        r = e.asset.members[0]
        row = [i, e.time.local.date().isoformat() if e.time.local else '', e.time.local.isoformat() if e.time.local else '',
               utc_text(e.time.utc) if e.time.utc else '', e.time.raw, e.time.source, e.time.note,
               r.path, e.asset.kind, r.gps.latitude, r.gps.longitude, r.gps.status,
               '跳过' if e.reason else '已加入行程', e.reason, e.asset.motion, '\n'.join(m.path for m in e.asset.members[1:])]
        append(all_rows, row)
        if e.reason:
            append(skipped, row)
    with path.open('xb') as f:
        book.save(f)


def export_journeys(result, parent, offset='+08:00'):
    entries, days = prepare_journeys(result, offset)
    parent = Path(parent).absolute()
    if not parent.is_dir():
        raise ValueError('请选择存在的导出目录')
    name = '旅行行程_' + datetime.now().strftime('%Y%m%d_%H%M%S')
    # Always create a fresh directory; never overwrite user files.
    for n in range(1000):
        output = parent / (name + (f'_{n}' if n else ''))
        try:
            output.mkdir()
            break
        except FileExistsError:
            continue
    else:
        raise FileExistsError('无法创建新的行程输出目录')
    try:
        for day, points in days.items():
            with (output / (day + '行程.kml')).open('xb') as f:
                f.write(kml_bytes(day, points, result.cancelled))
        write_workbook(result, entries, days, offset, output / '拍摄顺序与跳过清单.xlsx')
        (output / '使用说明.txt').write_text(
            f'每个 KML 是一天的行程，可导入支持 KML 的地图软件。\n行程时区：UTC{offset}\n'
            f'生成 {len(days)} 天、{sum(len(v) for v in days.values())} 个地点；跳过 {sum(bool(e.reason) for e in entries)} 项。\n'
            + ('本次扫描已取消，行程仅包含已读取部分。\n' if result.cancelled else '')
            + '仅一个地点的日期只生成地点标记，不画路线。\n同一天的拍摄地点按顺序连线，不是道路导航或连续 GPS 记录。\n'
            + '请核对视频时间和境外旅行的时区；默认中国时间 UTC+08:00。\n'
            + 'KML 无照片本身；Excel 中可查看所有条目的时间、路径、坐标及跳过原因。\n', encoding='utf-8-sig')
    except BaseException:
        shutil.rmtree(output)  # Only the exclusively created output directory, never source media.
        raise
    return str(output), len(days), sum(len(v) for v in days.values()), sum(bool(e.reason) for e in entries)
