import hashlib
import os
from pathlib import Path
import struct
import tempfile
import threading
import unittest
from unittest.mock import patch

from media_gps.core import *
from media_gps.exiftool import ExifTool, Cancelled
from media_gps.export import export_xlsx
from openpyxl import load_workbook


def meta(**kw):
    return {'File:FileType': 'JPEG', **kw}


def record(path, kind=PHOTO, ident='', status=COMPLETE):
    return Record(str(path), kind, 'MP4' if kind == VIDEO else 'JPEG', '', GPS(status, 30., 120., '测试'), identifier=ident)


def mp4():
    return b''.join(struct.pack('>I4s', 8 + len(v), k) + v for k, v in [(b'ftyp', b'isom\x00\x00\x00\x00'), (b'moov', b''), (b'mdat', b'\x00' * 16)])


class GPSTests(unittest.TestCase):
    def test_missing_only_on_success(self):
        self.assertEqual(gps_from_metadata(meta()).status, MISSING)
        for m, error in [(meta(), '读取失败'), (meta(**{'ExifTool:Warning': 'Truncated'}), ''), ({}, ''), ({'File:FileType': 'TXT'}, '')]:
            self.assertEqual(gps_from_metadata(m, error).status, UNKNOWN)

    def test_signed_and_refs(self):
        g = gps_from_metadata(meta(**{'GPS:GPSLatitude': 30, 'GPS:GPSLongitude': 120, 'GPS:GPSLatitudeRef': 'S', 'GPS:GPSLongitudeRef': 'W'}))
        self.assertEqual((g.status, g.latitude, g.longitude), (COMPLETE, -30, -120))
        g = gps_from_metadata(meta(**{'Composite:GPSLatitude': -20.5, 'Composite:GPSLongitude': 100.1}))
        self.assertEqual(g.status, COMPLETE)

    def test_exif_refs_required(self):
        self.assertEqual(gps_from_metadata(meta(**{'GPS:GPSLatitude': 30, 'GPS:GPSLongitude': 120})).status, INVALID)

    def test_partial_ranges_and_nonfinite(self):
        for fields in [
            {'Composite:GPSLatitude': 30},
            {'Composite:GPSLatitude': 91, 'Composite:GPSLongitude': 120},
            {'Composite:GPSLatitude': 30, 'Composite:GPSLongitude': -181},
            {'Composite:GPSLatitude': 'NaN', 'Composite:GPSLongitude': 10},
            {'Composite:GPSLatitude': float('inf'), 'Composite:GPSLongitude': 10},
            {'Composite:GPSLatitude': True, 'Composite:GPSLongitude': 10},
            {'GPS:GPSLatitude': -30, 'GPS:GPSLongitude': 120, 'GPS:GPSLatitudeRef': 'N', 'GPS:GPSLongitudeRef': 'E'},
        ]:
            with self.subTest(fields=fields):
                self.assertEqual(gps_from_metadata(meta(**fields)).status, INVALID)

    def test_zero_zero_flagged_but_equator_allowed(self):
        self.assertEqual(gps_from_metadata(meta(**{'Composite:GPSPosition': '0 0'})).status, INVALID)
        self.assertEqual(gps_from_metadata(meta(**{'Composite:GPSPosition': '0 120'})).status, COMPLETE)

    def test_coordinate_sources_conflict(self):
        self.assertEqual(gps_from_metadata(meta(**{'Composite:GPSPosition': '30 120', 'XMP:GPSPosition': '31 120'})).status, INVALID)

    def test_coordinate_formats(self):
        for v in ('+30.2500-120.5000+100.0/', '30.25 -120.5 100', [30.25, -120.5, 100]):
            g = gps_from_metadata({'File:FileType': 'MP4', 'Keys:GPSCoordinates': v})
            self.assertEqual((g.status, g.latitude, g.longitude), (COMPLETE, 30.25, -120.5))

    def test_unparsed_gps_hint_not_missing(self):
        self.assertEqual(gps_from_metadata(meta(**{'Vendor:GPSLog': 'opaque'})).status, UNKNOWN)

    def test_never_mix_sources(self):
        self.assertEqual(gps_from_metadata(meta(**{'A:GPSLatitude': 30, 'B:GPSLongitude': 120})).status, INVALID)


class MotionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / '动态照片.jpg'
        self.video = mp4()
        self.path.write_bytes(b'\xff\xd8\xffx\xff\xd9' + self.video)

    def test_offset_structure_and_readonly(self):
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        r = make_record(self.path, meta(**{'XMP-GCamera:MicroVideo': 1, 'XMP-GCamera:MicroVideoOffset': len(self.video)}))
        self.assertEqual(r.kind, MOTION)
        self.assertTrue(r.motion.startswith('发现'))
        self.assertEqual(before, hashlib.sha256(self.path.read_bytes()).hexdigest())

    def test_container_structure(self):
        m = meta(**{'XMP-GContainer:Directory': [{'Item': {'Semantic': 'Primary', 'Mime': 'image/jpeg'}}, {'Item': {'Semantic': 'MotionPhoto', 'Mime': 'video/mp4', 'Length': len(self.video)}}]})
        self.assertTrue(make_record(self.path, m).motion.startswith('发现'))

    def test_stale_marker_is_uncertain(self):
        r = make_record(self.path, meta(**{'XMP-GCamera:MotionPhoto': 1}))
        self.assertEqual(r.kind, MOTION)
        self.assertIn('无法确认', r.motion)

    def test_bad_offset_or_truncated_video(self):
        for offset in (999999, 2, 'nonsense'):
            r = make_record(self.path, meta(**{'XMP-GCamera:MicroVideoOffset': offset}))
            self.assertFalse(r.motion.startswith('发现'))
        self.path.write_bytes(b'prefix' + self.video[:-4])
        self.assertFalse(validate_mp4_region(self.path, 6, len(self.video)-4)[0])

    def test_disabled_motion_marker(self):
        r = make_record(self.path, meta(**{'XMP-GCamera:MotionPhoto': 0}))
        self.assertEqual(r.kind, PHOTO)

    def test_plain_image_not_called_motion(self):
        self.assertEqual(make_record(self.path, meta()).kind, PHOTO)


class PairTests(unittest.TestCase):
    def test_unique_id_deduplicated_and_gps_not_borrowed(self):
        p = record('/x/a.jpg', ident='pair', status=MISSING)
        v = record('/x/b.mov', VIDEO, 'pair')
        a = group_records([p, v], SCOPES[0])
        self.assertEqual(len(a), 1)
        self.assertEqual(a[0].gps_status, MISSING)
        self.assertEqual(len(a[0].members), 2)
        self.assertEqual(group_records([p, v], SCOPES[3]), [])
        self.assertEqual(len(group_records([p, v], SCOPES[2])), 1)

    def test_ambiguous_identifiers_not_merged(self):
        r = [record('/x/a.jpg', ident='same'), record('/x/b.jpg', ident='same'), record('/x/a.mov', VIDEO, 'same')]
        self.assertEqual(len(group_records(r, SCOPES[0])), 3)

    def test_same_stem_only_candidate(self):
        a = group_records([record('/x/a.jpg'), record('/x/a.mp4', VIDEO)], SCOPES[0])
        self.assertEqual(len(a), 2)
        self.assertIn('疑似', a[0].motion)
        self.assertTrue(a[0].problem)

    def test_scopes(self):
        r = [record('/x/a.jpg'), record('/x/b.mp4', VIDEO), record('/x/c.heic', MOTION)]
        self.assertEqual([len(group_records(r, s)) for s in SCOPES], [3, 1, 1, 1])


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_recursive_unknown_extension_unicode_and_skips(self):
        sub = self.root / '小米 相册 #💡'
        sub.mkdir()
        (sub / '无扩展名').write_bytes(b'\xff\xd8\xffabcdef')
        (sub / 'bad.heic').write_bytes(b'broken')
        (sub / 'note.txt').write_text('not media')
        found, skipped = discover(self.root, threading.Event(), lambda *a: None, [])
        self.assertEqual(len(found), 2)
        self.assertEqual(skipped, 1)

    def test_link_not_followed(self):
        if os.name == 'nt':
            self.skipTest('Windows 创建符号链接需要额外权限')
        (self.root / 'a.jpg').write_bytes(b'jpg')
        (self.root / 'loop').symlink_to(self.root, target_is_directory=True)
        (self.root / 'link.jpg').symlink_to(self.root / 'a.jpg')
        notices = []
        found, _ = discover(self.root, threading.Event(), lambda *a: None, notices)
        self.assertEqual(len(found), 1)
        self.assertEqual(len(notices), 2)

    def test_scan_export_readonly_and_formula_safety(self):
        p = self.root / '中文.jpg'
        p.write_bytes(b'fixture')
        before = p.read_bytes(), p.stat().st_mtime_ns
        class Reader:
            version = 'test'
            def read(self, paths, cancel):
                return {str(p): (meta(**{'EXIF:DateTimeOriginal': '=1+1'}), '') for p in paths}
        result = scan(self.root, Reader())
        out = self.root / 'report.xlsx'
        export_xlsx(result, out)
        self.assertEqual((p.read_bytes(), p.stat().st_mtime_ns), before)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ['扫描说明', '媒体结果', '成员明细'])
        self.assertEqual(wb['媒体结果']['E2'].data_type, 's')
        self.assertEqual(wb['媒体结果']['H2'].value, MISSING)
        wb.close()
        with self.assertRaises(FileExistsError):
            export_xlsx(result, out)

    def test_cancel_preserves_partial_records(self):
        for i in range(33):
            (self.root / f'{i:03}.jpg').write_bytes(b'fixture')
        event = threading.Event()
        class Reader:
            def read(self, paths, cancel):
                return {str(p): (meta(), '') for p in paths}
        def progress(phase, done, total):
            if phase == '读取元数据' and done >= 16:
                event.set()
        result = scan(self.root, Reader(), event, progress)
        self.assertTrue(result.cancelled)
        self.assertEqual(len(result.records), 16)
        self.assertEqual(result.discovered, 33)

    def test_thousand_files_batch_bound(self):
        for i in range(1200):
            (self.root / f'{i}.jpg').write_bytes(b'fixture')
        batches = []
        class Reader:
            def read(self, paths, cancel):
                batches.append(len(paths))
                return {str(p): (meta(), '') for p in paths}
        result = scan(self.root, Reader())
        self.assertEqual(len(result.assets), 1200)
        self.assertLessEqual(max(batches), 16)

    def test_missing_response_and_parse_failure_unknown(self):
        (self.root / 'bad.jpg').write_bytes(b'broken')
        class Reader:
            def read(self, paths, cancel):
                return {}
        result = scan(self.root, Reader())
        self.assertEqual(result.records[0].gps.status, UNKNOWN)


class ReaderTests(unittest.TestCase):
    def reader(self):
        with patch('media_gps.exiftool.locate_exiftool', return_value='exiftool'):
            return ExifTool()

    def test_invalid_json_unknown(self):
        r = self.reader()
        p = Path('/tmp/a.jpg').absolute()
        with patch.object(r, '_run', return_value=(0, b'not json', b'')):
            output = r.read([p], threading.Event())
        self.assertIn('解析失败', output[str(p)][1])

    def test_line_injection_rejected(self):
        r = self.reader()
        p = Path('/tmp/a\n-overwrite_original\n.jpg').absolute()
        with patch.object(r, '_run') as run:
            output = r.read([p], threading.Event())
            run.assert_not_called()
        self.assertIn('无法安全', output[str(p)][1])

    def test_command_is_readonly_and_utf8(self):
        r = self.reader()
        p = Path('/tmp/小米 #照片.jpg').absolute()
        with patch.object(r, '_run', return_value=(0, b'[]', b'')) as run:
            r.read([p], threading.Event())
            args, data, _ = run.call_args.args
            self.assertEqual(data.decode('utf-8'), str(p) + '\n')
            self.assertIn('filename=UTF8', args)
            self.assertFalse(any(a in args for a in ('-o', '-w', '-overwrite_original', '-tagsFromFile')))

    def test_timeout_and_stderr(self):
        r = self.reader()
        p = Path('/tmp/a.jpg').absolute()
        with patch.object(r, '_run', side_effect=TimeoutError('timeout')):
            self.assertIn('timeout', r.read([p], threading.Event())[str(p)][1])
        with patch.object(r, '_run', return_value=(0, b'[]', b'warning')):
            self.assertIn('stderr', r.read([p], threading.Event())[str(p)][1])

if __name__ == '__main__':
    unittest.main()
