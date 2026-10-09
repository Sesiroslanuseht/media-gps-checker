import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

try:
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication
    from media_gps.gui import Window
    HAS_QT = True
except ImportError:
    HAS_QT = False

from media_gps.core import *


@unittest.skipUnless(HAS_QT, '需要安装 PySide6')
class GuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.window = Window()
        self.window.show()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.window.folder.setText(str(self.root))

    def tearDown(self):
        if self.window.worker and self.window.worker.isRunning():
            self.window.worker.cancel.set()
            self.wait(lambda: not self.window.worker.isRunning())
        if self.window.export_worker and self.window.export_worker.isRunning():
            self.wait(lambda: not self.window.export_worker.isRunning())
        self.window.close()
        self.app.processEvents()

    def wait(self, condition, timeout=8):
        start = time.monotonic()
        while not condition():
            self.app.processEvents()
            if time.monotonic() - start > timeout:
                self.fail('GUI test timeout')
            time.sleep(0.005)
        self.app.processEvents()

    def test_background_scan_responsive_and_cancel(self):
        for i in range(160):
            (self.root / f'{i}.jpg').write_bytes(b'fixture')
        class Reader:
            version = 'fake'
            def __init__(self, *a):
                pass
            def check(self, *a):
                pass
            def read(self, paths, cancel):
                time.sleep(0.03)
                return {str(p): ({'File:FileType': 'JPEG'}, '') for p in paths}
        ticks = []
        timer = QTimer()
        timer.timeout.connect(lambda: ticks.append(True))
        timer.start(5)
        with patch('media_gps.gui.ExifTool', Reader):
            self.window.start_scan()
            self.wait(lambda: len(ticks) >= 8)
            self.assertFalse(self.window.start.isEnabled())
            self.window.cancel_scan()
            self.wait(lambda: self.window.result is not None and not self.window.worker.isRunning())
        timer.stop()
        self.assertTrue(self.window.result.cancelled)
        self.assertTrue(self.window.start.isEnabled())
        self.assertTrue(self.window.export.isEnabled())

    def test_filter_detail_and_background_export(self):
        complete = Record(str(self.root/'a.jpg'), PHOTO, 'JPEG', '', GPS(COMPLETE, 30, 120, '通过'))
        missing = Record(str(self.root/'b.jpg'), PHOTO, 'JPEG', '', GPS(MISSING, reason='未找到'))
        r = ScanResult(assets=[Asset([complete], PHOTO), Asset([missing], PHOTO)], records=[complete, missing], discovered=2)
        self.window.on_result(r)
        self.assertEqual(self.window.model.rowCount(), 2)
        self.window.filter.setChecked(True)
        self.assertEqual(self.window.model.rowCount(), 1)
        self.window.table.selectRow(0)
        self.app.processEvents()
        self.assertIn('b.jpg', self.window.detail.toPlainText())
        path = self.root/'report.xlsx'
        with patch('media_gps.gui.QFileDialog.getSaveFileName', return_value=(str(path), '')):
            self.window.export_result()
            self.wait(lambda: not self.window.export_worker.isRunning())
        from openpyxl import load_workbook
        book = load_workbook(path)
        self.assertEqual(book['媒体结果'].max_row, 3)  # Filter does not silently exclude rows from export.
        book.close()

    def test_journey_button_sort_and_daily_export(self):
        from test_journey import asset
        r = ScanResult(assets=[asset('late.jpg','2026:10:04 18:00:00'), asset('early.mp4','2026:10:04 08:00:00',VIDEO), asset('next.jpg','2026:10:05 08:00:00')], discovered=3)
        self.window.on_result(r)
        self.window.set_busy(False)
        self.assertTrue(self.window.journeys.isEnabled())
        self.assertEqual(self.window.model.assets[0].members[0].path,'early.mp4')
        self.assertIn('08:00:00', self.window.model.data(self.window.model.index(0,1)))
        self.window.filter.setChecked(True)  # Complete points are hidden, but still exported.
        self.assertEqual(self.window.model.rowCount(),0)
        with patch('media_gps.gui.QFileDialog.getExistingDirectory',return_value=str(self.root)):
            self.window.export_journey()
            self.wait(lambda: not self.window.export_worker.isRunning())
        dirs=list(self.root.glob('旅行行程_*'))
        self.assertEqual(len(dirs),1)
        self.assertEqual(len(list(dirs[0].glob('*.kml'))),2)
        self.assertTrue(self.window.journeys.isEnabled())
        self.assertIn('3 个地点', self.window.status.text())
