from datetime import datetime, timezone
from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET
from unittest.mock import patch

from openpyxl import load_workbook
from media_gps.core import *
from media_gps.capture_time import capture_time, fixed_timezone
from media_gps.journey import prepare_journeys, export_journeys, KML


def asset(name, taken='', kind=PHOTO, status=COMPLETE, meta=None, lat=30.25, lon=120.5):
    r = Record(name, kind, 'MP4' if kind == VIDEO else 'JPEG', taken, GPS(status, lat, lon, '夹具'), metadata=meta or {})
    return Asset([r], kind)


class TimeTests(unittest.TestCase):
    def test_sort_uses_capture_not_filename(self):
        r = ScanResult(assets=[asset('a.jpg','2026:10:04 14:00:00'), asset('z.mp4','2026:10:04 09:00:00',VIDEO)])
        entries, _ = prepare_journeys(r)
        self.assertEqual([e.asset.members[0].path for e in entries], ['z.mp4', 'a.jpg'])

    def test_explicit_timezone_converts_before_day_group(self):
        a = asset('a.jpg', '2026-10-04T23:30:00Z')
        t = capture_time(a.members[0])
        self.assertEqual(t.local.date().isoformat(), '2026-10-05')
        self.assertEqual(t.utc.hour, 23)

    def test_exif_offset_and_subsecond(self):
        a = asset('a.jpg', meta={'ExifIFD:DateTimeOriginal':'2026:10:04 08:00:00', 'ExifIFD:OffsetTimeOriginal':'+09:00', 'ExifIFD:SubSecTimeOriginal':'125'})
        t = capture_time(a.members[0])
        self.assertEqual((t.local.hour,t.local.microsecond), (7,125000))

    def test_quicktime_standard_utc(self):
        a = asset('a.mp4', kind=VIDEO, meta={'QuickTime:CreateDate':'2026:10:04 01:00:00'})
        self.assertEqual(capture_time(a.members[0]).local.hour,9)
        self.assertIn('QuickTime', capture_time(a.members[0]).note)

    def test_no_timezone_visible_assumption(self):
        a = asset('a.jpg','2026:10:04 08:00:00')
        t = capture_time(a.members[0], '+09:00')
        self.assertEqual(t.local.hour,8)
        self.assertIn('+09:00',t.note)

    def test_invalid_missing_and_conflicting_time(self):
        for meta in [
            {'ExifIFD:DateTimeOriginal':'0000:00:00 00:00:00'},
            {'ExifIFD:DateTimeOriginal':'2026:02:30 12:00:00'},
            {'ExifIFD:DateTimeOriginal':'2026:10:04'},
            {'System:FileModifyDate':'2026:10:04 08:00:00'},
            {'A:DateTimeOriginal':'2026:10:04 08:00:00','B:DateTimeOriginal':'2026:10:04 09:00:00'},
        ]:
            with self.subTest(meta=meta):
                self.assertIsNone(capture_time(asset('a.jpg',meta=meta).members[0]).utc)

    def test_timezone_validation(self):
        for bad in ('china','+15:00','+08:60','+14:30','8'):
            with self.assertRaises(ValueError): fixed_timezone(bad)
        fixed_timezone('+05:30')

    def test_same_time_stable_tie_and_unknown_last(self):
        r=ScanResult(assets=[asset('z.jpg','2026:10:04 08:00:00'),asset('0.jpg'),asset('a.jpg','2026:10:04 08:00:00')])
        entries,_=prepare_journeys(r)
        self.assertEqual([e.asset.members[0].path for e in entries],['a.jpg','z.jpg','0.jpg'])


class JourneyTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)

    def test_daily_xml_order_single_point_and_readonly(self):
        source=self.root/'原照片.jpg'
        source.write_bytes(b'original content')
        before=(source.read_bytes(),source.stat().st_mtime_ns)
        r=ScanResult(assets=[asset(str(source),'2026:10:04 12:00:00',lon=121),asset('先拍<&.mp4','2026:10:04 09:00:00',VIDEO),asset('next.jpg','2026:10:05 08:00:00')])
        out,days,points,skipped=export_journeys(r,self.root)
        self.assertEqual((days,points,skipped),(2,3,0))
        ns={'k':KML}
        day1=ET.parse(Path(out)/'2026-10-04行程.kml')
        self.assertEqual(day1.getroot().tag,'{'+KML+'}kml')
        line=day1.find('.//k:LineString/k:coordinates',ns).text.split()
        self.assertEqual(line,['120.50000000,30.25000000,0','121.00000000,30.25000000,0'])
        names=[m.find('k:name',ns).text for m in day1.findall('.//k:Placemark',ns) if m.find('k:Point',ns) is not None]
        self.assertIn('先拍<&.mp4',names[0])
        self.assertEqual(day1.find('.//k:TimeStamp/k:when',ns).text,'2026-10-04T01:00:00Z')
        day2=ET.parse(Path(out)/'2026-10-05行程.kml')
        self.assertIsNone(day2.find('.//k:LineString',ns))
        self.assertEqual((source.read_bytes(),source.stat().st_mtime_ns),before)

    def test_skip_reasons_and_workbook_chronology(self):
        r=ScanResult(assets=[asset('badgps.jpg','2026:10:04 08:00:00',status=MISSING),asset('badtime.jpg'),asset('valid.jpg','2026:10:04 09:00:00')])
        out,days,points,skipped=export_journeys(r,self.root)
        self.assertEqual((days,points,skipped),(1,1,2))
        book=load_workbook(Path(out)/'拍摄顺序与跳过清单.xlsx')
        rows=list(book['拍摄顺序清单'].values)
        self.assertEqual([row[7] for row in rows[1:]],['badgps.jpg','valid.jpg','badtime.jpg'])
        self.assertEqual(book['跳过项'].max_row,3)
        self.assertIn(MISSING,book['跳过项']['N2'].value)
        book.close()

    def test_paired_photo_only_once_and_no_borrowing(self):
        p=asset('a.jpg','2026:10:04 08:00:00')
        p.kind=MOTION
        p.members.append(asset('a.mov','2026:10:04 08:00:00',VIDEO,status=MISSING).members[0])
        entries,days=prepare_journeys(ScanResult(assets=[p]))
        self.assertEqual(sum(map(len,days.values())),1)
        p.members[0].gps.status=MISSING
        p.members[1].gps.status=COMPLETE
        _,days=prepare_journeys(ScanResult(assets=[p]))
        self.assertEqual(days,{})

    def test_cancelled_scan_label_no_valid_points_and_no_overwrite(self):
        r=ScanResult(assets=[asset('a.jpg')],cancelled=True)
        a,*counts=export_journeys(r,self.root)
        b,*_=export_journeys(r,self.root)
        self.assertNotEqual(a,b)
        self.assertEqual(counts,[0,0,1])
        self.assertIn('已取消',(Path(a)/'使用说明.txt').read_text(encoding='utf-8-sig'))
        self.assertEqual(list(Path(a).glob('*.kml')),[])

    def test_non_wgs84_and_defensive_coordinates(self):
        for a in [asset('a.jpg','2026:10:04 08:00:00',meta={'GPS:GPSMapDatum':'TOKYO'}),asset('b.jpg','2026:10:04 08:00:00',lat=float('nan'))]:
            entries,days=prepare_journeys(ScanResult(assets=[a]))
            self.assertEqual(days,{})
            self.assertTrue(entries[0].reason)

    def test_failed_export_cleans_only_new_output(self):
        keep=self.root/'existing.txt'
        keep.write_text('keep')
        with patch('media_gps.journey.write_workbook',side_effect=OSError('disk full')):
            with self.assertRaises(OSError):
                export_journeys(ScanResult(assets=[asset('a.jpg','2026:10:04 08:00:00')]),self.root)
        self.assertEqual(list(self.root.iterdir()),[keep])
